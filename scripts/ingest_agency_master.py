"""
ingest_agency_master.py
------------------------
Loads `data/uploads/Agency Master Data.XLSX` into `sceis_agency_master`.

The source xlsx has 4,645 rows representing every (Cost_Center,
Functional_Area, AGY_Funded_Program) combination. The canonical
table is keyed on Cost_Center alone (one row per cost center), so
this loader DEDUPES to one row per Cost_Center, picking the entry
with the latest Valid_From date as the representative row. The
non-key fields (Functional_Area, AGY_Funded_Program, etc.) are
therefore lossy — the canonical purpose of this table is to expose
Cost_Center → Name and the most recent classification metadata.

If a downstream consumer needs the full many-to-many breakdown,
the schema needs to change to a composite PK (Cost_Center,
Functional_Area, AGY_Funded_Program) and the loader needs to keep
all rows.

IDEMPOTENCY: TRUNCATE-then-INSERT.

Usage:
    python scripts/ingest_agency_master.py [--dry-run]
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path

import duckdb
import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DB_PATH = REPO_ROOT / "db" / "scde.duckdb"
SOURCE_PATH = REPO_ROOT / "data" / "uploads" / "Agency Master Data.XLSX"


def to_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_rows(src: Path) -> list[dict]:
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    header = list(next(iterator))
    idx = {h: i for i, h in enumerate(header)}

    required = [
        "Cost Center", "Name", "Functional Area", "Functional Area Description",
        "Mini Code", "State Funded Program", "AGY Funded Program",
        "AGY Funded Program Name", "Valid from", "Valid to",
    ]
    missing = [c for c in required if c not in idx]
    if missing:
        raise SystemExit(f"Missing columns in {src.name}: {missing}")

    # Dedupe: keep the row with the latest Valid_From per Cost Center.
    by_cc: dict[str, dict] = {}
    for r in iterator:
        if r is None:
            continue
        cc = to_str(r[idx["Cost Center"]])
        if not cc:
            continue
        valid_from = to_date(r[idx["Valid from"]])
        candidate = {
            "Cost_Center":             cc,
            "Name":                    to_str(r[idx["Name"]]),
            "Functional_Area":         to_str(r[idx["Functional Area"]]),
            "Functional_Area_Desc":    to_str(r[idx["Functional Area Description"]]),
            "Mini_Code":               to_str(r[idx["Mini Code"]]),
            "State_Funded_Program":    to_str(r[idx["State Funded Program"]]),
            "AGY_Funded_Program":      to_str(r[idx["AGY Funded Program"]]),
            "AGY_Funded_Program_Name": to_str(r[idx["AGY Funded Program Name"]]),
            "Valid_From":              valid_from,
            "Valid_To":                to_date(r[idx["Valid to"]]),
        }
        existing = by_cc.get(cc)
        if existing is None:
            by_cc[cc] = candidate
        else:
            # Prefer the row with the later Valid_From; fall back to first seen.
            ev = existing.get("Valid_From")
            if valid_from is not None and (ev is None or valid_from > ev):
                by_cc[cc] = candidate
    return list(by_cc.values())


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="parse only; do not write")
    args = ap.parse_args()

    if not SOURCE_PATH.exists():
        raise SystemExit(f"Source missing: {SOURCE_PATH}")

    rows = parse_rows(SOURCE_PATH)
    print(f"Source: {SOURCE_PATH.name}")
    print(f"Distinct cost centers: {len(rows)}")
    print(f"Sample: {rows[0]}")

    if args.dry_run:
        print("DRY RUN — not writing.")
        return

    cols = list(rows[0].keys())
    placeholders = ",".join("?" * len(cols))
    sql = f"INSERT INTO sceis_agency_master ({','.join(cols)}) VALUES ({placeholders})"
    con = duckdb.connect(str(DB_PATH))
    con.execute("DELETE FROM sceis_agency_master")
    con.executemany(sql, [[r[c] for c in cols] for r in rows])
    after = con.execute("SELECT COUNT(*) FROM sceis_agency_master").fetchone()[0]
    print(f"Loaded {after} rows into sceis_agency_master")
    con.close()


if __name__ == "__main__":
    main()
