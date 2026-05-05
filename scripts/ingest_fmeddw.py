"""
ingest_fmeddw.py
-----------------
Ingestion script for SAP Funds Management Drilldown Reporting (FMEDDW)
xlsx exports into the `sceis_fmeddw` table in db/scde.duckdb.

IDEMPOTENCY STRATEGY: TRUNCATE-then-INSERT.
  Each FMEDDW extract is single-FY. Re-running always produces a clean
  load from the current `*FMEDDW*.xlsx` file. To merge multiple FYs, this
  script must be revised — currently it truncates the table.

SOURCE FILE: data/uploads/*FMEDDW*.xlsx
  Picks the most recently modified file matching the glob. Single sheet,
  single-row header, native numeric amounts, no merged cells.

SOURCE → SCHEMA COLUMN MAPPING (20 source columns → 22 schema columns):
  The xlsx has TWO columns named 'Document Status' — string ('Posted' /
  'Preposted Posted') and integer (1 / 3). The integer is renamed to
  Document_Status_Code at load. Two derived columns are added:
    Is_Rollup     – TRUE when LEN(Funds_Center) = 8 (FM hierarchy parent)
    Source_File   – originating filename for provenance

SIGN CONVENTION (per internal-budget agent):
  Send entries are NEGATIVE; Receive / Enter / Supplement are POSITIVE.
  The view vw_budget_vs_actuals_by_funds_center relies on this — never
  abs() before aggregation.

Usage:
    python scripts/ingest_fmeddw.py [--dry-run]

    --dry-run  Parse and report counts without writing to the database.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path

import duckdb
import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DB_PATH = REPO_ROOT / "db" / "scde.duckdb"
UPLOADS_DIR = REPO_ROOT / "data" / "uploads"


def find_source_file() -> Path:
    candidates = sorted(
        UPLOADS_DIR.glob("*FMEDDW*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise SystemExit(f"No FMEDDW xlsx found in {UPLOADS_DIR}")
    return candidates[0]


def to_int(v):
    if v is None or v == "":
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v)
    s = str(v).strip()
    return int(s) if s else None


def to_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unparseable date: {v!r}")


def to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    return float(s) if s else None


def parse_rows(src: Path):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    header = list(next(iterator))

    # Two 'Document Status' columns. First (string) keeps the name; second
    # is renamed to Document_Status_Code for disambiguation.
    seen_doc_status = False
    norm_header = []
    for h in header:
        if h == "Document Status":
            if not seen_doc_status:
                norm_header.append("Document_Status")
                seen_doc_status = True
            else:
                norm_header.append("Document_Status_Code")
        else:
            norm_header.append(h)
    idx = {col: i for i, col in enumerate(norm_header)}

    required = [
        "Entry Document", "Entry Document Line", "Document Date", "Document Year",
        "Version", "Entry Document Type", "Process", "Created on", "Fiscal Year",
        "Budget Type", "Fund", "Funds Center", "Commitment Item", "Functional Area",
        "Grant", "Document_Status", "Funded Program",
        "Total of Transactions in Local Currency", "Entry currency",
        "Document_Status_Code",
    ]
    missing = [c for c in required if c not in idx]
    if missing:
        raise SystemExit(f"Missing expected columns in {src.name}: {missing}")

    rows = []
    for r in iterator:
        if r is None or all(c is None for c in r):
            continue
        entry_doc = to_str(r[idx["Entry Document"]])
        entry_line = to_str(r[idx["Entry Document Line"]])
        if not entry_doc or not entry_line:
            continue
        funds_center = to_str(r[idx["Funds Center"]])
        rows.append({
            "Entry_Document":       entry_doc,
            "Entry_Document_Line":  entry_line,
            "Document_Date":        to_date(r[idx["Document Date"]]),
            "Document_Year":        to_str(r[idx["Document Year"]]),
            "Version":              to_str(r[idx["Version"]]),
            "Entry_Document_Type":  to_str(r[idx["Entry Document Type"]]),
            "Process":              to_str(r[idx["Process"]]),
            "Created_On":           to_date(r[idx["Created on"]]),
            "Fiscal_Year":          to_int(r[idx["Fiscal Year"]]),
            "Budget_Type":          to_str(r[idx["Budget Type"]]),
            "Fund":                 to_str(r[idx["Fund"]]),
            "Funds_Center":         funds_center,
            "Commitment_Item":      to_str(r[idx["Commitment Item"]]),
            "Functional_Area":      to_str(r[idx["Functional Area"]]),
            "Grant":                to_str(r[idx["Grant"]]),
            "Document_Status":      to_str(r[idx["Document_Status"]]),
            "Document_Status_Code": to_int(r[idx["Document_Status_Code"]]),
            "Funded_Program":       to_str(r[idx["Funded Program"]]),
            "Amount":               to_float(r[idx["Total of Transactions in Local Currency"]]),
            "Currency":             to_str(r[idx["Entry currency"]]),
            "Is_Rollup":            (len(funds_center) == 8) if funds_center else False,
            "Source_File":          src.name,
        })
    return rows


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="parse only; do not write to DB")
    args = ap.parse_args()

    src = find_source_file()
    print(f"Source: {src.name}")

    rows = parse_rows(src)
    print(f"Parsed {len(rows):,} rows")

    fy_counts = {}
    for r in rows:
        fy_counts[r["Fiscal_Year"]] = fy_counts.get(r["Fiscal_Year"], 0) + 1
    print("By Fiscal_Year:", dict(sorted(fy_counts.items())))

    rollup_count = sum(1 for r in rows if r["Is_Rollup"])
    print(f"Is_Rollup rows: {rollup_count} (filter out for per-center aggregation)")

    if args.dry_run:
        print("DRY RUN — not writing to DB.")
        return

    cols = list(rows[0].keys())
    placeholders = ",".join("?" * len(cols))
    insert_sql = f"INSERT INTO sceis_fmeddw ({','.join(cols)}) VALUES ({placeholders})"

    con = duckdb.connect(str(DB_PATH))
    con.execute("DELETE FROM sceis_fmeddw")
    con.executemany(insert_sql, [[r[c] for c in cols] for r in rows])
    after = con.execute("SELECT COUNT(*) FROM sceis_fmeddw").fetchone()[0]
    print(f"Loaded {after:,} rows into sceis_fmeddw")

    # Spot-check view
    sample = con.execute("""
        SELECT Funds_Center, Total_Budget, Actuals, Available, Pct_Consumed
        FROM vw_budget_vs_actuals_by_funds_center
        WHERE Funds_Center='H630HG0010'
    """).fetchall()
    print("Spot-check H630HG0010:", sample)

    con.close()
    print("Done.")


if __name__ == "__main__":
    main()
