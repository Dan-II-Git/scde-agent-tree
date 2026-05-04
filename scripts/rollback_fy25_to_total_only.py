"""EMERGENCY ROLLBACK — restore FY25 lea_wpu_category to TOTAL-only state.

Reverses the categorical FY25 load committed on 2026-05-04 (the 675-row
detailed-ADM load). Re-derives the 75 TOTAL rows from WPU 13525.xlsx,
which is the same source the engine consumed before the detailed-ADM
work began. Idempotent: safe to run more than once.

Default: applies the rollback immediately (this is an emergency tool;
optimized for one-command paste during a presentation crisis).
Pass --dry-run to preview without writing.

Run:
  python scripts/rollback_fy25_to_total_only.py            # apply rollback
  python scripts/rollback_fy25_to_total_only.py --dry-run  # preview only

After rollback, the categorical detail can be restored at any time with:
  python scripts/load_lea_wpu_category_fy25_adm.py --commit
"""
from __future__ import annotations

import argparse
from pathlib import Path

import duckdb
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "db" / "scde.duckdb"
SOURCE = ROOT / "data" / "uploads" / "WPU 13525.xlsx"  # space in name is intentional

FISCAL_YEAR = 2025
REPORT_CYCLE = 135


def parse_total_rows(path: Path) -> list[dict]:
    """Read the 'WPU' sheet's two-column district-total layout."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "WPU" not in wb.sheetnames:
        raise RuntimeError(f"No 'WPU' sheet in {path.name}")
    ws = wb["WPU"]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[1:]:  # skip header
        raw_dist = row[0] if len(row) > 0 else None
        raw_wpu  = row[1] if len(row) > 1 else None
        if not raw_dist or raw_wpu in (None, ""):
            continue
        out.append({
            "district_name_raw": str(raw_dist).strip(),
            "wpu":               float(raw_wpu),
        })
    return out


def resolve_district_ids(con, records: list[dict]) -> list[dict]:
    dim = {n.strip(): did for did, n in
           con.execute("SELECT District_ID, long_name FROM dim_district").fetchall()}
    out = []
    for r in records:
        did = dim.get(r["district_name_raw"])
        if did is None:
            print(f"  WARN unresolved district name: {r['district_name_raw']!r}")
            continue
        out.append({**r, "district_id": did})
    return out


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--dry-run", action="store_true",
                   help="Preview only. Default behavior applies the rollback.")
    args = p.parse_args()

    if not SOURCE.exists():
        raise SystemExit(f"Source not found: {SOURCE}")

    print(f"ROLLBACK: FY{FISCAL_YEAR} lea_wpu_category -> TOTAL-only")
    print(f"Source: {SOURCE.name}")

    records = parse_total_rows(SOURCE)
    print(f"Parsed {len(records)} district rows from WPU sheet")

    con = duckdb.connect(str(DB_PATH), read_only=args.dry_run)
    matched = resolve_district_ids(con, records)
    print(f"Resolved {len(matched)} to District_IDs ({len(records) - len(matched)} dropped)")

    pre_n = con.execute(
        "SELECT COUNT(*) FROM lea_wpu_category WHERE Fiscal_Year = ? AND Report_Cycle = ?",
        [FISCAL_YEAR, REPORT_CYCLE],
    ).fetchone()[0]
    pre_cats = sorted(r[0] for r in con.execute(
        "SELECT DISTINCT Category FROM lea_wpu_category WHERE Fiscal_Year = ? AND Report_Cycle = ?",
        [FISCAL_YEAR, REPORT_CYCLE],
    ).fetchall())
    print(f"Current state : {pre_n} rows · categories={pre_cats}")
    print(f"Target state  : {len(matched)} rows · categories=['TOTAL']")

    if args.dry_run:
        print("\n[dry run] no write performed.")
        return

    con.execute("BEGIN")
    try:
        con.execute(
            "DELETE FROM lea_wpu_category WHERE Fiscal_Year = ? AND Report_Cycle = ?",
            [FISCAL_YEAR, REPORT_CYCLE],
        )
        con.executemany(
            """INSERT INTO lea_wpu_category
               (District_ID, Fiscal_Year, Report_Cycle, Category,
                ADM, Weighted_Pupils, Source_Row_Labels)
               VALUES (?, ?, ?, 'TOTAL', ?, ?, 'FY25 135-day WPU (total only) [rollback]')""",
            [(r["district_id"], FISCAL_YEAR, REPORT_CYCLE, r["wpu"], r["wpu"])
             for r in matched],
        )
        con.execute("COMMIT")
    except Exception:
        con.execute("ROLLBACK")
        raise

    post_n = con.execute(
        "SELECT COUNT(*) FROM lea_wpu_category WHERE Fiscal_Year = ? AND Report_Cycle = ?",
        [FISCAL_YEAR, REPORT_CYCLE],
    ).fetchone()[0]
    print(f"\nROLLED BACK: {pre_n} -> {post_n} rows in lea_wpu_category "
          f"(FY={FISCAL_YEAR}, Report_Cycle={REPORT_CYCLE})")
    print("Re-apply the categorical load any time with:")
    print("  python scripts/load_lea_wpu_category_fy25_adm.py --commit")


if __name__ == "__main__":
    main()
